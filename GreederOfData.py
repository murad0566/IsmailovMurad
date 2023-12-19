import numpy as np
import pandas as pd
import math as mt
import os
import datetime

from datetime import datetime as dt
from datetime import timedelta
import warnings
import shutil
import time
import seaborn as sns

warnings.filterwarnings("ignore", category = FutureWarning)
pd.set_option('display.max_columns', 100)

from adlib import PrettyWriter as pw
from adlib.BruteFormater import BruteData as bf
from adlib import LoadMaster as lm
from adlib import Counter as cnt

from optbinning import BinningProcess

from sklearn.linear_model import LogisticRegression
from sklearn.ensemble import GradientBoostingClassifier
from sklearn.metrics import classification_report
from sklearn.metrics import auc, roc_auc_score, roc_curve
from sklearn.model_selection import train_test_split
from sklearn.pipeline import Pipeline
from sklearn.feature_selection import RFE

from statsmodels.stats.outliers_influence import variance_inflation_factor
from matplotlib import pyplot as plt
from tqdm import tqdm
from patsy import dmatrices
from random import shuffle

import statsmodels.api as sm
import openpyxl
from itertools import chain
import gc
import lightgbm as lgb
import seaborn as sns

import pickle
import hashlib
hash_f = hashlib.blake2b(digest_size = 10)


def write_user_hist(f):
    path_hist = r'D:\Обмен\_Скрипты\Резервные копии рабочих папок\Папка Директора по автоматизации/hist/' + str(datetime.datetime.now().date()) +'_'+ os.getlogin() + '.pkl'
    try:
        try:
            with open(path_hist, 'rb') as fp:
                t = pickle.load(fp)
        except:
            t = {}

        hash_f.update(bytes(str(time.time()), 'utf-8'))
        t.update({hash_f.hexdigest():{'user':os.getlogin(),'time':str(datetime.datetime.now()),'modul':'GreederOfData','func':f}})
        with open(path_hist, 'wb') as fp:
            pickle.dump(t, fp)
    except:
        pass

class Modeler(object):
    """
    Мастер объект модельера. Позволяет осуществить полный цикл построения модели,
    начиная нахождением оптимальных категорий переменных, заканчивая сохранением модели.
    
    Пример подбора модели:
    >>> model = god.Modeler(d, 'Target')
    >>> model.bin_data(min_hit_rate = 0.1)

    >>> model.feature_greeder(max_feature_allow = 10, rounds = 5, max_pvalue_allow = 0.015)
    >>> best_features = model.fr.sort_values('gini', ascending = False).Features.iloc[0]

    >>> model.logit_model(best_features)
    >>> model.save_model('./model.json')
    """

    def __init__(self, data: pd.DataFrame, target: str):
        """
        Инициализация класса модельера. Необходимо предедать DataFrame, и название целевой переменной.
        В целевой переменной не должно быть пропусков. Признаки должны быть числового формата.
        """
        write_user_hist('Modeler')
        self.d = data.replace([-np.inf, np.inf], np.nan).select_dtypes(include = 'number').astype(float)
        self.d = self.d.reset_index(drop = True)
        if type(target) == str:
            self.target = data[target]
            self.target_name = target
        else:
            self.target = target
            self.target_name = target.name 
            try:
                self.d = self.d.join(self.target, how = 'left')
            except:
                pass
        self.d = self.d[self.d[self.target_name].notna()]
        self.max_pvalue_allow = 0.015
        self.max_feature_allow = 5
        self.max_vif_allow = 1.5
        self.constant = True
        self.must_feature = []
        pass
   
    def bin_data(self,
        cols_to_bin: list = None,
        min_n_bins: int = None,
        max_n_bins: int = None,
        min_bin_size: float = 0.1,
        min_hit_rate: float = 0.2,
        must_event_rate: float = False,
        missing_merge_perc: float = 0.0) -> None: 
        """
        Оптимальне разбиение переменных на категории. Данные разбитые на категории сохраняются внутри модуля.
        Чтобы посмотреть итоговую таблицу разбитых переменных - нужно прописать: 
        >>> model.summ
        
        Чтобы посмотреть как разбилась конкретная переменная - нужно прописать, где Feature1 - имя переменной: 
        >>> optd['Feature1']
        
        Параметры:
        cols_to_bin - список из название переменных, которые нужно разбить на категории. Если ничего не передается - то разбиваются все переменные.
        min_n_bins - минимально допустимое количество категорий в разбитой переменной. Если ничего не передается - то разбивается оптимально.
        max_n_bins - максимальное допустимое количество категорий в разбитой переменной. Если ничего не передается - то разбивается оптимально.
        min_bin_size - минимальная доля категории в разбивке. Если ничего не передается - то минимальная доля ставится 10%.
        min_hit_rate - минимальная доля валидных (не пропуском) значений в переменной. Если доля валидных значений меньше заданного значения - переменная исключается из итоговой выборки. Если ничего не передается - то минималная доля валидных ставится 20%.
        must_event_rate - доля произошедшего события.  Если ничего не передается - то доля сохраняется исходная.
        missing_merge_perc - доля пропишенных значений в переменной меньше которой, категория пропусков объединяется с ближайшей по WoE.
        """
        write_user_hist('bin_data')
        def get_test_sample(data, target, must_event_rate = False):
            if must_event_rate:
                target_name = target.name
                target_count = target.sum()
                current_event_rate = target_count / data.shape[0]
                if current_event_rate < must_event_rate:
                    to_add_sample = int(target_count / must_event_rate) - int(target_count)

                    t = data.join(target)
                    return t[t[target_name] == 1].append(t[t[target_name] == 0].sample(to_add_sample)).reset_index(drop = True)
                else:
                    to_add_sample = int((data.shape[0] - target_count) * must_event_rate / (1 - must_event_rate))

                    t = data.join(target)
                    return t[t[target_name] == 0].append(t[t[target_name] == 1].sample(to_add_sample)).reset_index(drop = True)

            else:
                return data.join(target).reset_index(drop = True)
        
        def get_new_missing_cat(col):
            subject = self.optd[col]
            subject = subject[subject.Count != 0]
            woes = subject[subject.Bin.str.contains(',')].WoE
            if len(woes) == 1:
                return 'pass'
            missing_woe = subject[subject.Bin == 'Missing'].WoE.values

            abss = list(abs(woes - missing_woe))
            return abss.index(min(abss))


        self.cols_to_bin = cols_to_bin
        self.min_n_bins = min_n_bins
        self.max_n_bins = max_n_bins
        self.min_bin_size = min_bin_size
        self.min_hit_rate = min_hit_rate
        self.must_event_rate = must_event_rate
        self.missing_merge_perc = missing_merge_perc
        self.summ, self.optd, self.cat_df, self.bp = self.get_bins(cols_to_bin, min_n_bins, max_n_bins, min_bin_size, min_hit_rate)
        self.test = self.bp.transform(self.d, metric  = 'indices', metric_missing = -1)
        if self.missin_merge_perc != 0:
            to_miss_merge = (self.d.isna().sum() / self.d.shape[0])
            to_miss_merge_list = list(to_miss_merge[(to_miss_merge <= missing_merge_perc) & (to_miss_merge != 0)].index)
            for col in to_miss_merge_list:
                to_replace_missting = get_new_missing_cat(col)
                if to_replace_missting == 'pass':
                    continue
                self.test[col] = self.test[col].replace(-1, to_replace_missting)
        
        self.t = get_test_sample(self.test, self.d[self.target_name], must_event_rate)
        self.target = self.t[self.target_name]
        
    def get_params(self, args):
        write_user_hist('get_params')
        to_return = []
        if 'opt_info' == args:
            print('Optimal Binning Info')
            print(' cols_to_bin: ', self.cols_to_bin, '\n',
                    'min_n_bins: ', self.min_n_bins, '\n',
                    'max_n_bins: ', self.max_n_bins, '\n',
                    'min_bin_size: ', self.min_bin_size, '\n',
                    'min_hit_rate: ', self.min_hit_rate)
            return
        if 'best_feature' == args:
            return self.fr
        if 'summury' in args:
            to_return.append(self.summ)
        if 'optd' in args:
            to_return.append(self.optd)
        if 'cat_df' in args:
            to_return.append(self.cat_df)       
        if 'bp' in args:
            to_return.append(self.bp)
        if len(to_return) == 0:
            print('Args not found')
            return
        if type(args) == str:
            return to_return[0]
        return to_return
            
    def get_bins(self, cols_to_bin, min_n_bins, max_n_bins, min_bin_size, min_hit_rate):
        write_user_hist('get_bins')
        if min_hit_rate:
            hit_table = (self.d.count() / self.d.shape[0]).to_frame()
            d = self.d[hit_table[hit_table[0] >= min_hit_rate].index]
            hit_table = hit_table.reset_index() 
            hit_table.columns = ['name', 'hit_rate']   
        else:
            hit_table = (self.d.count() / self.d.shape[0]).to_frame()
            d = self.d
            hit_table = hit_table.reset_index() 
            hit_table.columns = ['name', 'hit_rate']
            
        if cols_to_bin:
            x = d[cols_to_bin].values
        else:
            
            x = d.drop(columns = [self.target_name]).values
            cols_to_bin = list(d.columns)
            cols_to_bin.remove(self.target_name)

        y = d[self.target_name].values
        binning_process = BinningProcess(variable_names = cols_to_bin,
                                         min_n_bins = min_n_bins, max_n_bins = max_n_bins,
                                         min_bin_size = min_bin_size, n_jobs = -1, max_pvalue_policy = 'all')
        binning_process.fit(x, y)

        summary_res = binning_process.summary().merge(hit_table, how = 'left', on = 'name')
        cols_to_corr = summary_res.name.values
        opt_ditc = {}
        for col in cols_to_corr:
            opt_ditc[col] = binning_process.get_binned_variable(col).binning_table.build()

        new_x = binning_process.transform(x)
        new_d = pd.DataFrame(data = new_x, columns = cols_to_bin)
        return summary_res, opt_ditc, new_d, binning_process
 
    def feature_greeder(self, max_pvalue_allow = 0.015, by_score = 'all', max_feature_allow = 5, max_vif_allow = 1.5, constant = True, must_feature = [], top = 5000, rounds = 5):
        write_user_hist('feature_greeder')
        self.max_pvalue_allow = max_pvalue_allow
        self.max_feature_allow = max_feature_allow
        self.max_vif_allow = max_vif_allow
        self.by_score = by_score
        self.constant = constant
        self.must_feature = must_feature
        warnings.filterwarnings("ignore", message = 'Maximum Likelihood optimization failed to ')
        
        
        def get_params(best_feature):
            def get_to_drop_col(col):
                return to_fit.filter(like = col).sum().to_frame().sort_values(by = 0).index[-1]

            y = self.target
            target_name = self.target_name
            res_feature = self.must_feature
            best_bic = 0
            max_pvalue_allow = self.max_pvalue_allow
            max_feature_allow = self.max_feature_allow
            max_vif_allow = self.max_vif_allow
            constant = self.constant
            first_best_feat = True
                
            for feature in best_feature:
                col_to_train = []
                col_to_train = res_feature.copy()
                col_to_train.append(feature)
                to_fit = pd.get_dummies(self.t[col_to_train], columns = col_to_train, prefix = col_to_train)
                to_drop = list(map(get_to_drop_col, col_to_train))
                to_fit.drop(columns = to_drop, inplace = True)

                if constant:
                    to_fit = sm.add_constant(to_fit)
                log_model = sm.Logit(y, to_fit, maxiter = 35)

                try:
                    res_model = log_model.fit(disp=0)
                except:
                    continue


                form = target_name+' ~ ' + '+'.join(col_to_train)
                y, X = dmatrices(form, data = self.t[col_to_train].join(self.t[target_name]), return_type = 'dataframe')
                vifs = [variance_inflation_factor(X.values, i) for i in range(X.shape[1])][1:]


                if (res_model.pvalues.max() <=  max_pvalue_allow) & (~res_model.pvalues.isna().any()) & (best_bic < res_model.bic) & (max(vifs) <= max_vif_allow):
                    res_feature = col_to_train.copy()
                    if first_best_feat == True:
                        first_best_feat = col_to_train[-1]
                if len(res_feature) == max_feature_allow:
                    break

            to_fit = pd.get_dummies(self.t[res_feature], columns = res_feature, prefix = res_feature)
            to_drop = list(map(get_to_drop_col, res_feature))
            to_fit.drop(columns = to_drop, inplace = True)
            if constant:
                to_fit = sm.add_constant(to_fit)

            log_model = sm.Logit(y, to_fit, maxiter = 35)
            log_model = log_model.fit(disp=0)

            probs = log_model.predict(to_fit)
            fpr1, tpr1, threshold = roc_curve(y, probs)
            roc_auc = auc(fpr1, tpr1)

            return {'Features':res_feature, 'roc':roc_auc, 'gini':roc_auc*2-1}, first_best_feat
      
        def load_to_model(best_feature):
            feature_spliter = 0
            best_feature = list(best_feature)
            for i in tqdm(range(rounds)):
                
                try:
                    dict_to_ap, prev_best_feature = get_params(best_feature[feature_spliter:])
                    feature_spliter = best_feature.index(prev_best_feature) + 1
                    dict_to_ap.update({'by_score':score_type})
                    best_res.append(dict_to_ap)
                except Exception as e:
                    print(e)
        
        if by_score == 'all':
            all_score = ['iv', 'js', 'gini', 'quality_score']
        else:
            if type(by_score) == str:
                all_score = [by_score]
            else:
                all_score = by_score
            
        best_res = []
        for score_type in all_score:
            load_to_model(self.summ.sort_values(score_type, ascending = False).head(top).name.values)

        self.fr = pd.DataFrame(best_res)
        self.fr['n_feature'] = self.fr.Features.apply(lambda x: len(x))
        self.fr = self.fr.sort_values("gini", ascending = False).reset_index(drop = True)   
        
    def logit_model(self, col_to_train, constant = True, print_info = True):
        write_user_hist('logit_model')
        def get_to_drop_col(col):
            return to_fit.filter(like = col).sum().to_frame().sort_values(by = 0).index[-1]
        self.col_to_train = col_to_train
        to_fit = pd.get_dummies(self.t[col_to_train], columns = col_to_train, prefix = col_to_train)
        to_drop = list(map(get_to_drop_col, col_to_train))
        to_fit.drop(columns = to_drop, inplace = True)
        if constant:
            to_fit = sm.add_constant(to_fit)
        
        log_model = sm.Logit(self.target, to_fit, maxiter = 35)
        self.res_model = log_model.fit()
        if print_info:
            print(self.res_model.summary())
        
        probs = self.res_model.predict(to_fit)
        fpr1, tpr1, threshold = roc_curve(self.target, probs)
        self.roc_auc = auc(fpr1, tpr1)

        if print_info:
            print('Roc: ',self.roc_auc)
            self.gini = self.roc_auc * 2 - 1
            print('Gini: ',self.gini)

            plt.plot(fpr1, tpr1)
            plt.plot([0,1],[0,1],'k--')
        
    def save_model(self, path):
        write_user_hist('save_model')
        def get_bins_cat(col_to_train):
            if col_to_train != 'const':
                bins = [-np.inf] + list(self.bp.get_binned_variable(col_to_train).splits) + [np.inf]

                to_fit = pd.get_dummies(self.t[col_to_train], columns = col_to_train, prefix = col_to_train)
                #to_drop = get_to_drop_col(col_to_train)

                counts = to_fit.sum().to_frame()
                counts = counts.reset_index()
                counts.columns = ['categries','counts']
                if counts.categries.values[0][-2:] == '-1':
                    nan_ex = True
                else:
                    nan_ex = False


                params = self.res_model.params.to_frame()
                params = params.reset_index()
                params.columns = ['categries', 'params']
                params = counts.merge(params, on = 'categries', how = 'left')
                params['params'] = params['params'].fillna(0)
                if nan_ex:
                    return {col_to_train : {'nan':params.iloc[0].params, 'bins':bins, 'coefs': list(params.iloc[1:].params.values)}}
                return {col_to_train : {'bins':bins, 'coefs': list(params.params.values)}}

            else:
                to_const = pd.DataFrame( [{'categries':'const', 'counts':'1'}])
                params = self.res_model.params.to_frame()
                params = params.reset_index()
                params.columns = ['categries', 'params']
                params = to_const.merge(params, on = 'categries', how = 'left')
                params['params'] = params['params'].fillna(0)
                return {col_to_train : params.iloc[0].params}
                
        const_exist = False
        if 'const' in self.res_model.params.index:
            const_exist = True
            
        cols = self.col_to_train
        model_info = {}
        for col in cols:
            model_info.update(get_bins_cat(col))
        if const_exist:
            model_info.update(get_bins_cat('const'))
            
        lm.save(model_info, path)
        
    def save_info(self, path, sheet_name = 'Model'):
        write_user_hist('save_info')
        pic_name = 'roc.png'
        row_cnt = 1
        s1 = pd.DataFrame(self.res_model.summary().tables[0].data)
        s2 = pd.DataFrame(self.res_model.summary().tables[1].data)

        pw.write_df(path, sheet_name, s1, write_cell = 'A'+str(row_cnt), write_headers=False, write_index=False, borders = True)
        row_cnt+=s1.shape[0]+2
        pw.write_df(path, sheet_name, s2, write_cell = 'A'+str(row_cnt), write_headers=False, write_index=False, borders = True)
        row_cnt+=s2.shape[0]+2



        form = self.target_name + ' ~ ' + '+'.join(self.col_to_train)
        y, X = dmatrices(form, data = self.t[self.col_to_train].join(self.t[self.target.name]), return_type = 'dataframe')
        vifs = [variance_inflation_factor(X.values, i) for i in range(X.shape[1])]
        vif = pd.DataFrame()
        vif['name'] = [self.target_name] + list(self.col_to_train)
        vif['VIF'] = vifs
        summ1 = self.summ.merge(vif, on = 'name', how = 'left')
        summ1 = summ1[summ1.name.isin(self.col_to_train)]

        pw.write_df(path, sheet_name, summ1, write_cell = 'A'+str(row_cnt), write_headers=True, write_index=False, borders = True)
        row_cnt+=summ1.shape[0]+2



        probs = self.res_model.predict()
        fpr1, tpr1, threshold = roc_curve(self.target, probs)
        roc_auc = auc(fpr1, tpr1)
        pw.write_cell(path, sheet_name, 'ROC:', write_cell = 'A'+str(row_cnt),font_size = 14 ,font_bold = True)
        pw.write_cell(path, sheet_name, roc_auc, write_cell = 'B'+str(row_cnt))
        row_cnt+=1
        pw.write_cell(path, sheet_name, 'Gini:', write_cell = 'A'+str(row_cnt),font_size = 14 ,font_bold = True)
        pw.write_cell(path, sheet_name, str(roc_auc*2-1), write_cell = 'B'+str(row_cnt))
        row_cnt+=2

        plt.plot(fpr1, tpr1)
        plt.plot([0,1],[0,1],'k--')
        plt.savefig(pic_name, bbox_inches = 'tight', dpi = 100)

        # вставка графика 
        wb = openpyxl.load_workbook(path)
        ws = wb[sheet_name]
        img_n = openpyxl.drawing.image.Image(pic_name)
        img_n.anchor = 'C'+str(row_cnt) 
        ws.add_image(img_n)
        wb.save(path)
        row_cnt+=20
        os.remove(pic_name)

        for col in self.col_to_train:
            to_w = self.optd[col]

            pw.write_cell(path, sheet_name, col, write_cell = 'A'+str(row_cnt),font_size = 14 ,font_bold = True)
            row_cnt+=1
            pw.write_df(path, sheet_name, to_w, write_cell = 'A'+str(row_cnt), write_headers=True, write_index=False)
            row_cnt+=to_w.shape[0]+2

class FeatureSelector():
    """
    Class for performing feature selection for machine learning or data preprocessing.
    
    Implements five different methods to identify features for removal 
    
        1. Find columns with a missing percentage greater than a specified threshold
        2. Find columns with a single unique value
        3. Find collinear variables with a correlation greater than a specified correlation coefficient
        4. Find features with 0.0 feature importance from a gradient boosting machine (gbm)
        5. Find low importance features that do not contribute to a specified cumulative feature importance from the gbm
        
    Parameters
    --------
        data : dataframe
            A dataset with observations in the rows and features in the columns
        labels : array or series, default = None
            Array of labels for training the machine learning model to find feature importances. These can be either binary labels
            (if task is 'classification') or continuous targets (if task is 'regression').
            If no labels are provided, then the feature importance based methods are not available.
        
    Attributes
    --------
    
    ops : dict
        Dictionary of operations run and features identified for removal
        
    missing_stats : dataframe
        The fraction of missing values for all features
    
    record_missing : dataframe
        The fraction of missing values for features with missing fraction above threshold
        
    unique_stats : dataframe
        Number of unique values for all features
    
    record_single_unique : dataframe
        Records the features that have a single unique value
        
    corr_matrix : dataframe
        All correlations between all features in the data
    
    record_collinear : dataframe
        Records the pairs of collinear variables with a correlation coefficient above the threshold
        
    feature_importances : dataframe
        All feature importances from the gradient boosting machine
    
    record_zero_importance : dataframe
        Records the zero importance features in the data according to the gbm
    
    record_low_importance : dataframe
        Records the lowest importance features not needed to reach the threshold of cumulative importance according to the gbm
    
    
    Notes
    --------
    
        - All 5 operations can be run with the `identify_all` method.
        - If using feature importances, one-hot encoding is used for categorical variables which creates new columns
    
    """
    
    def __init__(self, data, labels=None):
        write_user_hist('FeatureSelector')
        # Dataset and optional training labels
        try:
            self.data = data.drop(columns = [labels.name])
        except:
            self.data = data
        self.labels = labels

        if labels is None:
            print('No labels provided. Feature importance based methods are not available.')
        
        self.base_features = list(data.columns)
        self.one_hot_features = None
        
        # Dataframes recording information about features to remove
        self.record_missing = None
        self.record_single_unique = None
        self.record_collinear = None
        self.record_zero_importance = None
        self.record_low_importance = None
        
        self.missing_stats = None
        self.unique_stats = None
        self.corr_matrix = None
        self.feature_importances = None
        
        # Dictionary to hold removal operations
        self.ops = {}
        
        self.one_hot_correlated = False
        
    def identify_missing(self, missing_threshold):
        """Find the features with a fraction of missing values above `missing_threshold`"""
        
        self.missing_threshold = missing_threshold

        # Calculate the fraction of missing in each column 
        missing_series = self.data.isnull().sum() / self.data.shape[0]
        self.missing_stats = pd.DataFrame(missing_series).rename(columns = {'index': 'feature', 0: 'missing_fraction'})

        # Sort with highest number of missing values on top
        self.missing_stats = self.missing_stats.sort_values('missing_fraction', ascending = False)

        # Find the columns with a missing percentage above the threshold
        record_missing = pd.DataFrame(missing_series[missing_series > missing_threshold]).reset_index().rename(columns = 
                                                                                                               {'index': 'feature', 
                                                                                                                0: 'missing_fraction'})

        to_drop = list(record_missing['feature'])

        self.record_missing = record_missing
        self.ops['missing'] = to_drop
        
        print('%d features with greater than %0.2f missing values.\n' % (len(self.ops['missing']), self.missing_threshold))
        
    def identify_single_unique(self):
        """Finds features with only a single unique value. NaNs do not count as a unique value. """

        # Calculate the unique counts in each column
        unique_counts = self.data.nunique()
        self.unique_stats = pd.DataFrame(unique_counts).rename(columns = {'index': 'feature', 0: 'nunique'})
        self.unique_stats = self.unique_stats.sort_values('nunique', ascending = True)
        
        # Find the columns with only one unique count
        record_single_unique = pd.DataFrame(unique_counts[unique_counts == 1]).reset_index().rename(columns = {'index': 'feature', 
                                                                                                                0: 'nunique'})

        to_drop = list(record_single_unique['feature'])
    
        self.record_single_unique = record_single_unique
        self.ops['single_unique'] = to_drop
        
        print('%d features with a single unique value.\n' % len(self.ops['single_unique']))
    
    def identify_collinear(self, correlation_threshold, one_hot=False):
        """
        Finds collinear features based on the correlation coefficient between features. 
        For each pair of features with a correlation coefficient greather than `correlation_threshold`,
        only one of the pair is identified for removal. 
        Using code adapted from: https://chrisalbon.com/machine_learning/feature_selection/drop_highly_correlated_features/
        
        Parameters
        --------
        correlation_threshold : float between 0 and 1
            Value of the Pearson correlation cofficient for identifying correlation features
        one_hot : boolean, default = False
            Whether to one-hot encode the features before calculating the correlation coefficients
        """
        
        self.correlation_threshold = correlation_threshold
        self.one_hot_correlated = one_hot
        
         # Calculate the correlations between every column
        if one_hot:
            
            # One hot encoding
            features = pd.get_dummies(self.data)
            self.one_hot_features = [column for column in features.columns if column not in self.base_features]

            # Add one hot encoded data to original data
            self.data_all = pd.concat([features[self.one_hot_features], self.data], axis = 1)
            
            corr_matrix = pd.get_dummies(features).corr()

        else:
            corr_matrix = self.data.corr()
        
        self.corr_matrix = corr_matrix
    
        # Extract the upper triangle of the correlation matrix
        upper = corr_matrix.where(np.triu(np.ones(corr_matrix.shape), k = 1).astype(np.bool))
        
        # Select the features with correlations above the threshold
        # Need to use the absolute value
        to_drop = [column for column in upper.columns if any(upper[column].abs() > correlation_threshold)]

        # Dataframe to hold correlated pairs
        record_collinear = pd.DataFrame(columns = ['drop_feature', 'corr_feature', 'corr_value'])

        # Iterate through the columns to drop to record pairs of correlated features
        for column in to_drop:

            # Find the correlated features
            corr_features = list(upper.index[upper[column].abs() > correlation_threshold])

            # Find the correlated values
            corr_values = list(upper[column][upper[column].abs() > correlation_threshold])
            drop_features = [column for _ in range(len(corr_features))]    

            # Record the information (need a temp df for now)
            temp_df = pd.DataFrame.from_dict({'drop_feature': drop_features,
                                             'corr_feature': corr_features,
                                             'corr_value': corr_values})

            # Add to dataframe
            record_collinear = record_collinear.append(temp_df, ignore_index = True)

        self.record_collinear = record_collinear
        self.ops['collinear'] = to_drop
        
        print('%d features with a correlation magnitude greater than %0.2f.\n' % (len(self.ops['collinear']), self.correlation_threshold))

    def identify_zero_importance(self, task, eval_metric=None, 
                                 n_iterations=10, early_stopping = True):
        """
        
        Identify the features with zero importance according to a gradient boosting machine.
        The gbm can be trained with early stopping using a validation set to prevent overfitting. 
        The feature importances are averaged over `n_iterations` to reduce variance. 
        
        Uses the LightGBM implementation (http://lightgbm.readthedocs.io/en/latest/index.html)
        Parameters 
        --------
        eval_metric : string
            Evaluation metric to use for the gradient boosting machine for early stopping. Must be
            provided if `early_stopping` is True
        task : string
            The machine learning task, either 'classification' or 'regression'
        n_iterations : int, default = 10
            Number of iterations to train the gradient boosting machine
            
        early_stopping : boolean, default = True
            Whether or not to use early stopping with a validation set when training
        
        
        Notes
        --------
        
        - Features are one-hot encoded to handle the categorical variables before training.
        - The gbm is not optimized for any particular task and might need some hyperparameter tuning
        - Feature importances, including zero importance features, can change across runs
        """

        if early_stopping and eval_metric is None:
            raise ValueError("""eval metric must be provided with early stopping. Examples include "auc" for classification or
                             "l2" for regression.""")
            
        if self.labels is None:
            raise ValueError("No training labels provided.")
        
        # One hot encoding
        features = pd.get_dummies(self.data)
        self.one_hot_features = [column for column in features.columns if column not in self.base_features]

        # Add one hot encoded data to original data
        self.data_all = pd.concat([features[self.one_hot_features], self.data], axis = 1)

        # Extract feature names
        feature_names = list(features.columns)

        # Convert to np array
        features = np.array(features)
        labels = np.array(self.labels).reshape((-1, ))

        # Empty array for feature importances
        feature_importance_values = np.zeros(len(feature_names))
        
        print('Training Gradient Boosting Model\n')
        
        # Iterate through each fold
        for _ in range(n_iterations):

            if task == 'classification':
                model = lgb.LGBMClassifier(n_estimators=1000, learning_rate = 0.05, verbose = -1)

            elif task == 'regression':
                model = lgb.LGBMRegressor(n_estimators=1000, learning_rate = 0.05, verbose = -1)

            else:
                raise ValueError('Task must be either "classification" or "regression"')
                
            # If training using early stopping need a validation set
            if early_stopping:
                
                train_features, valid_features, train_labels, valid_labels = train_test_split(features, labels, test_size = 0.15, stratify=labels)

                # Train the model with early stopping
                model.fit(train_features, train_labels, eval_metric = eval_metric,
                          eval_set = [(valid_features, valid_labels)],
                          early_stopping_rounds = 100, verbose = -1)
                
                # Clean up memory
                gc.enable()
                del train_features, train_labels, valid_features, valid_labels
                gc.collect()
                
            else:
                model.fit(features, labels)

            # Record the feature importances
            feature_importance_values += model.feature_importances_ / n_iterations

        feature_importances = pd.DataFrame({'feature': feature_names, 'importance': feature_importance_values})

        # Sort features according to importance
        feature_importances = feature_importances.sort_values('importance', ascending = False).reset_index(drop = True)

        # Normalize the feature importances to add up to one
        feature_importances['normalized_importance'] = feature_importances['importance'] / feature_importances['importance'].sum()
        feature_importances['cumulative_importance'] = np.cumsum(feature_importances['normalized_importance'])

        # Extract the features with zero importance
        record_zero_importance = feature_importances[feature_importances['importance'] == 0.0]
        
        to_drop = list(record_zero_importance['feature'])

        self.feature_importances = feature_importances
        self.record_zero_importance = record_zero_importance
        self.ops['zero_importance'] = to_drop
        
        print('\n%d features with zero importance after one-hot encoding.\n' % len(self.ops['zero_importance']))
    
    def identify_low_importance(self, cumulative_importance):
        """
        Finds the lowest importance features not needed to account for `cumulative_importance` fraction
        of the total feature importance from the gradient boosting machine. As an example, if cumulative
        importance is set to 0.95, this will retain only the most important features needed to 
        reach 95% of the total feature importance. The identified features are those not needed.
        Parameters
        --------
        cumulative_importance : float between 0 and 1
            The fraction of cumulative importance to account for 
        """

        self.cumulative_importance = cumulative_importance
        
        # The feature importances need to be calculated before running
        if self.feature_importances is None:
            raise NotImplementedError("""Feature importances have not yet been determined. 
                                         Call the `identify_zero_importance` method first.""")
            
        # Make sure most important features are on top
        self.feature_importances = self.feature_importances.sort_values('cumulative_importance')

        # Identify the features not needed to reach the cumulative_importance
        record_low_importance = self.feature_importances[self.feature_importances['cumulative_importance'] > cumulative_importance]

        to_drop = list(record_low_importance['feature'])

        self.record_low_importance = record_low_importance
        self.ops['low_importance'] = to_drop
    
        print('%d features required for cumulative importance of %0.2f after one hot encoding.' % (len(self.feature_importances) -
                                                                            len(self.record_low_importance), self.cumulative_importance))
        print('%d features do not contribute to cumulative importance of %0.2f.\n' % (len(self.ops['low_importance']),
                                                                                               self.cumulative_importance))
        
    def identify_all(self, selection_params):
        """
        Use all five of the methods to identify features to remove.
        
        Parameters
        --------
            
        selection_params : dict
           Parameters to use in the five feature selection methhods.
           Params must contain the keys ['missing_threshold', 'correlation_threshold', 'eval_metric', 'task', 'cumulative_importance']
        
        """
        
        # Check for all required parameters
        for param in ['missing_threshold', 'correlation_threshold', 'eval_metric', 'task', 'cumulative_importance']:
            if param not in selection_params.keys():
                raise ValueError('%s is a required parameter for this method.' % param)
        
        # Implement each of the five methods
        self.identify_missing(selection_params['missing_threshold'])
        self.identify_single_unique()
        self.identify_collinear(selection_params['correlation_threshold'])
        self.identify_zero_importance(task = selection_params['task'], eval_metric = selection_params['eval_metric'])
        self.identify_low_importance(selection_params['cumulative_importance'])
        
        # Find the number of features identified to drop
        self.all_identified = set(list(chain(*list(self.ops.values()))))
        self.n_identified = len(self.all_identified)
        
        print('%d total features out of %d identified for removal after one-hot encoding.\n' % (self.n_identified, 
                                                                                                  self.data_all.shape[1]))
        
    def check_removal(self, keep_one_hot=True):
        
        """Check the identified features before removal. Returns a list of the unique features identified."""
        
        self.all_identified = set(list(chain(*list(self.ops.values()))))
        print('Total of %d features identified for removal' % len(self.all_identified))
        
        if not keep_one_hot:
            if self.one_hot_features is None:
                print('Data has not been one-hot encoded')
            else:
                one_hot_to_remove = [x for x in self.one_hot_features if x not in self.all_identified]
                print('%d additional one-hot features can be removed' % len(one_hot_to_remove))
        
        return list(self.all_identified)
        
    
    def remove(self, methods, keep_one_hot = True):
        """
        Remove the features from the data according to the specified methods.
        
        Parameters
        --------
            methods : 'all' or list of methods
                If methods == 'all', any methods that have identified features will be used
                Otherwise, only the specified methods will be used.
                Can be one of ['missing', 'single_unique', 'collinear', 'zero_importance', 'low_importance']
            keep_one_hot : boolean, default = True
                Whether or not to keep one-hot encoded features
                
        Return
        --------
            data : dataframe
                Dataframe with identified features removed
                
        
        Notes 
        --------
            - If feature importances are used, the one-hot encoded columns will be added to the data (and then may be removed)
            - Check the features that will be removed before transforming data!
        
        """
        
        
        features_to_drop = []
      
        if methods == 'all':
            
            # Need to use one-hot encoded data as well
            data = self.data_all
                                          
            print('{} methods have been run\n'.format(list(self.ops.keys())))
            
            # Find the unique features to drop
            features_to_drop = set(list(chain(*list(self.ops.values()))))
            
        else:
            # Need to use one-hot encoded data as well
            if 'zero_importance' in methods or 'low_importance' in methods or self.one_hot_correlated:
                data = self.data_all
                
            else:
                data = self.data
                
            # Iterate through the specified methods
            for method in methods:
                
                # Check to make sure the method has been run
                if method not in self.ops.keys():
                    raise NotImplementedError('%s method has not been run' % method)
                    
                # Append the features identified for removal
                else:
                    features_to_drop.append(self.ops[method])
        
            # Find the unique features to drop
            features_to_drop = set(list(chain(*features_to_drop)))
            
        features_to_drop = list(features_to_drop)
            
        if not keep_one_hot:
            
            if self.one_hot_features is None:
                print('Data has not been one-hot encoded')
            else:
                             
                features_to_drop = list(set(features_to_drop) | set(self.one_hot_features))
       
        # Remove the features and return the data
        data = data.drop(columns = features_to_drop)
        self.removed_features = features_to_drop
        
        if not keep_one_hot:
        	print('Removed %d features including one-hot features.' % len(features_to_drop))
        else:
        	print('Removed %d features.' % len(features_to_drop))
        
        return data
    
    def plot_missing(self):
        """Histogram of missing fraction in each feature"""
        if self.record_missing is None:
            raise NotImplementedError("Missing values have not been calculated. Run `identify_missing`")
        
        self.reset_plot()
        
        # Histogram of missing values
        plt.style.use('seaborn-white')
        plt.figure(figsize = (7, 5))
        plt.hist(self.missing_stats['missing_fraction'], bins = np.linspace(0, 1, 11), edgecolor = 'k', color = 'red', linewidth = 1.5)
        plt.xticks(np.linspace(0, 1, 11));
        plt.xlabel('Missing Fraction', size = 14); plt.ylabel('Count of Features', size = 14); 
        plt.title("Fraction of Missing Values Histogram", size = 16);
        
    
    def plot_unique(self):
        """Histogram of number of unique values in each feature"""
        if self.record_single_unique is None:
            raise NotImplementedError('Unique values have not been calculated. Run `identify_single_unique`')
        
        self.reset_plot()

        # Histogram of number of unique values
        self.unique_stats.plot.hist(edgecolor = 'k', figsize = (7, 5))
        plt.ylabel('Frequency', size = 14); plt.xlabel('Unique Values', size = 14); 
        plt.title('Number of Unique Values Histogram', size = 16);
        
    
    def plot_collinear(self, plot_all = False):
        """
        Heatmap of the correlation values. If plot_all = True plots all the correlations otherwise
        plots only those features that have a correlation above the threshold
        
        Notes
        --------
            - Not all of the plotted correlations are above the threshold because this plots
            all the variables that have been idenfitied as having even one correlation above the threshold
            - The features on the x-axis are those that will be removed. The features on the y-axis
            are the correlated features with those on the x-axis
        
        Code adapted from https://seaborn.pydata.org/examples/many_pairwise_correlations.html
        """
        
        if self.record_collinear is None:
            raise NotImplementedError('Collinear features have not been idenfitied. Run `identify_collinear`.')
        
        if plot_all:
        	corr_matrix_plot = self.corr_matrix
        	title = 'All Correlations'
        
        else:
	        # Identify the correlations that were above the threshold
	        # columns (x-axis) are features to drop and rows (y_axis) are correlated pairs
	        corr_matrix_plot = self.corr_matrix.loc[list(set(self.record_collinear['corr_feature'])), 
	                                                list(set(self.record_collinear['drop_feature']))]

	        title = "Correlations Above Threshold"

       
        f, ax = plt.subplots(figsize=(10, 8))
        
        # Diverging colormap
        cmap = sns.diverging_palette(220, 10, as_cmap=True)

        # Draw the heatmap with a color bar
        sns.heatmap(corr_matrix_plot, cmap=cmap, center=0,
                    linewidths=.25, cbar_kws={"shrink": 0.6})

        # Set the ylabels 
        ax.set_yticks([x + 0.5 for x in list(range(corr_matrix_plot.shape[0]))])
        ax.set_yticklabels(list(corr_matrix_plot.index), size = int(160 / corr_matrix_plot.shape[0]));

        # Set the xlabels 
        ax.set_xticks([x + 0.5 for x in list(range(corr_matrix_plot.shape[1]))])
        ax.set_xticklabels(list(corr_matrix_plot.columns), size = int(160 / corr_matrix_plot.shape[1]));
        plt.title(title, size = 14)
        
    def plot_feature_importances(self, plot_n = 15, threshold = None):
        """
        Plots `plot_n` most important features and the cumulative importance of features.
        If `threshold` is provided, prints the number of features needed to reach `threshold` cumulative importance.
        Parameters
        --------
        
        plot_n : int, default = 15
            Number of most important features to plot. Defaults to 15 or the maximum number of features whichever is smaller
        
        threshold : float, between 0 and 1 default = None
            Threshold for printing information about cumulative importances
        """
        
        if self.record_zero_importance is None:
            raise NotImplementedError('Feature importances have not been determined. Run `idenfity_zero_importance`')
            
        # Need to adjust number of features if greater than the features in the data
        if plot_n > self.feature_importances.shape[0]:
            plot_n = self.feature_importances.shape[0] - 1

        self.reset_plot()
        
        # Make a horizontal bar chart of feature importances
        plt.figure(figsize = (10, 6))
        ax = plt.subplot()

        # Need to reverse the index to plot most important on top
        # There might be a more efficient method to accomplish this
        ax.barh(list(reversed(list(self.feature_importances.index[:plot_n]))), 
                self.feature_importances['normalized_importance'][:plot_n], 
                align = 'center', edgecolor = 'k')

        # Set the yticks and labels
        ax.set_yticks(list(reversed(list(self.feature_importances.index[:plot_n]))))
        ax.set_yticklabels(self.feature_importances['feature'][:plot_n], size = 12)

        # Plot labeling
        plt.xlabel('Normalized Importance', size = 16); plt.title('Feature Importances', size = 18)
        plt.show()

        # Cumulative importance plot
        plt.figure(figsize = (6, 4))
        plt.plot(list(range(1, len(self.feature_importances) + 1)), self.feature_importances['cumulative_importance'], 'r-')
        plt.xlabel('Number of Features', size = 14); plt.ylabel('Cumulative Importance', size = 14); 
        plt.title('Cumulative Feature Importance', size = 16);

        if threshold:

            # Index of minimum number of features needed for cumulative importance threshold
            # np.where returns the index so need to add 1 to have correct number
            importance_index = np.min(np.where(self.feature_importances['cumulative_importance'] > threshold))
            plt.vlines(x = importance_index + 1, ymin = 0, ymax = 1, linestyles='--', colors = 'blue')
            plt.show();

            print('%d features required for %0.2f of cumulative importance' % (importance_index + 1, threshold))

    def reset_plot(self):
        plt.rcParams = plt.rcParamsDefault

def build_model(data, model_info, suffix = ''):
    write_user_hist('build_model')
    X_list = []
    x_cnt = 1
    data['Z'+suffix] = 0
    
    for key in model_info:
        if key != 'const':
            data['X'+str(x_cnt)+suffix] =  pd.cut(data[key], bins = model_info[key]['bins'], labels = model_info[key]['coefs'], include_lowest = True).astype('float')
            try:
                data['X'+str(x_cnt)+suffix] = data['X'+str(x_cnt)+suffix].fillna(model_info[key]['nan'])
            except:
                pass

            data['Z'+suffix] += data['X'+str(x_cnt)+suffix]
            x_cnt+=1
        else:
            data['Z'+suffix] += model_info[key]
    
    def GetScore(Z):
        try:
            Score = round(1 / (1 + np.exp((-1) * Z)) * 100000, 2)
        except:
            Score = None
        return Score

    data['Score'+suffix+'_check'] = data['Z'+suffix].agg(GetScore) / 100000
    return data
    
def fds(d: pd.DataFrame) -> None:
    """
    Describe series or dataframe. Datetime objects presents like numeric.
    All numeric series describes with percentiles = [0.01, 0.05, 0.25, 0.5, 0.75, 0.95, 0.99].
    """


    write_user_hist('Full_Describe')
    try:
        hit = (d.count() / len(d)).to_frame().T
    except:
        d = d.to_frame()
        hit = (d.count() / len(d)).to_frame().T
    hit.index = ['valid']
    try:
        des_num = round(d.describe(percentiles = [0.01, 0.05, 0.25, 0.5, 0.75, 0.95, 0.99]), 2)
        des_num = des_num.append(hit)[des_num.columns]
    except:
        des_num = 'Numeric not found'
    try:
        des_date = d.describe(include = ['<M8[ns]'], datetime_is_numeric = True, percentiles = [0.01, 0.05, 0.25, 0.5, 0.75, 0.95, 0.99])
        des_date = des_date.append(hit)[des_date.columns]
    except:
        des_date = 'Datetime not found'
    try:
        des_obj = d.describe(include = ['O'])
        des_obj = des_obj.append(hit)[des_obj.columns]
    except:
        des_obj = 'Object not found'
    try:
        des_cat = d.describe(include = ['category'])
        des_cat = des_cat.append(hit)[des_cat.columns]
    except:
        des_cat = 'Categoty not found'
    display('Numeric dtypes', des_num, 'Datetime dtypes', des_date, 'Object dtypes', des_obj, 'Category dtypes', des_cat)    
    
def get_roc(x, y):
    write_user_hist('get_roc')
    fpr1, tpr1, threshold = roc_curve(x, y)
    roc_auc = auc(fpr1, tpr1)
    
    print('Roc: ',roc_auc)
    print('Gini: ',roc_auc*2 - 1)
    plt.plot(fpr1, tpr1)
    plt.plot([0,1],[0,1],'k--')